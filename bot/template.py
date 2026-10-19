"""Builds the provisional Excel template the owner fills in.

It holds headers, instructions, the bot's default messages, and the menu
structure described in CLAUDE.md. It contains no product, price, or policy data.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .knowledge import (
    DEFAULT_MESSAGES,
    FAQ_COLUMNS,
    MENU_ACTIONS,
    MENU_COLUMNS,
    MESSAGE_COLUMNS,
    PRODUCT_COLUMNS,
    REQUEST_COLUMNS,
    SHEET_FAQ,
    SHEET_MENU,
    SHEET_MESSAGES,
    SHEET_PRODUCTS,
    SHEET_REQUESTS,
)

HEADER_FILL = PatternFill("solid", fgColor="6B4F3A")
HEADER_FONT = Font(bold=True, color="FFFFFF")

STARTER_MENU = (
    ("1", "Cajas musicales", "productos", "Cajas musicales", "", "caja musical, cajas musicales, caja de musica, cajita musical"),
    ("2", "Otros productos", "productos", "Otros productos", "", "otros productos, catalogo, productos"),
    ("3", "Hablar con una persona", "persona", "", "", "persona, humano, asesor, asesora, cotizar, cotizacion, pedido, encargar"),
)

INSTRUCTIONS = (
    "Cómo usar este archivo",
    "",
    "El bot lee este Excel cada pocos minutos. No cambies los nombres de las hojas ni de las columnas (fila 1).",
    "Escribe solo información real del negocio: el bot no responde nada que no esté aquí.",
    "",
    "Productos: una fila por producto. Si Precio está vacío, el bot dice que la pieza se cotiza.",
    "  Categoría debe coincidir con la Categoría de la opción del menú que lista esos productos.",
    "  Disponible: escribe No para ocultar un producto sin borrarlo.",
    "  Palabras clave: palabras separadas por comas que hacen que el bot muestre ese producto.",
    "Preguntas: preguntas frecuentes (envíos, pagos, materiales, tiempos). Palabras clave separadas por comas.",
    "Menu: opciones numeradas del modo menú. Acción puede ser:",
    "  responder: envía el texto de la columna Respuesta.",
    "  productos: lista los productos de la Categoría indicada (vacía = todos).",
    "  persona: pasa la conversación a una persona y avisa por WhatsApp al número de administración.",
    "Mensajes: textos propios del bot. Puedes cambiar la columna Texto; no cambies la Clave.",
    "Solicitudes: la escribe el bot cada vez que alguien pide hablar con una persona. Puedes cambiar el Estado.",
    "",
    "Reglas para los textos: sin emojis y sin dos guiones seguidos.",
)


def _header(sheet: Worksheet, columns: tuple[str, ...], widths: tuple[int, ...]) -> None:
    sheet.append(list(columns))
    for index, cell in enumerate(sheet[1]):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        sheet.column_dimensions[cell.column_letter].width = widths[index] if index < len(widths) else 20
    sheet.freeze_panes = "A2"


def _wrap(sheet: Worksheet) -> None:
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_template() -> Workbook:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instrucciones"
    for line in INSTRUCTIONS:
        instructions.append([line])
    instructions["A1"].font = Font(bold=True, size=14)
    instructions.column_dimensions["A"].width = 110

    products = workbook.create_sheet(SHEET_PRODUCTS)
    _header(products, PRODUCT_COLUMNS, (28, 20, 50, 40, 14, 35, 12))
    available = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    products.add_data_validation(available)
    available.add("G2:G500")

    faq = workbook.create_sheet(SHEET_FAQ)
    _header(faq, FAQ_COLUMNS, (40, 70, 40))

    menu = workbook.create_sheet(SHEET_MENU)
    _header(menu, MENU_COLUMNS, (10, 30, 14, 20, 60, 50))
    for row in STARTER_MENU:
        menu.append(list(row))
    actions = DataValidation(type="list", formula1='"' + ",".join(MENU_ACTIONS) + '"', allow_blank=True)
    menu.add_data_validation(actions)
    actions.add("C2:C50")
    _wrap(menu)

    messages = workbook.create_sheet(SHEET_MESSAGES)
    _header(messages, MESSAGE_COLUMNS, (18, 70, 70))
    for key, (text, purpose) in DEFAULT_MESSAGES.items():
        messages.append([key, text, purpose])
    _wrap(messages)

    requests = workbook.create_sheet(SHEET_REQUESTS)
    _header(requests, REQUEST_COLUMNS, (18, 16, 24, 70, 14))
    return workbook
