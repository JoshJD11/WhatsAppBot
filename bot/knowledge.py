"""Business knowledge read from the Excel workbook.

The workbook layout is provisional (see README, "Excel layout"). Sheet and column
names are in Spanish because the owner edits them; the code matches them ignoring
case and accents.
"""

from dataclasses import dataclass, field
from io import BytesIO
from typing import Literal, cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .text import normalize

SHEET_PRODUCTS = "Productos"
SHEET_FAQ = "Preguntas"
SHEET_MENU = "Menu"
SHEET_MESSAGES = "Mensajes"
SHEET_REQUESTS = "Solicitudes"

PRODUCT_COLUMNS = ("Nombre", "Categoría", "Descripción", "Personalización", "Precio", "Palabras clave", "Disponible")
FAQ_COLUMNS = ("Pregunta", "Respuesta", "Palabras clave")
MENU_COLUMNS = ("Opción", "Texto del menú", "Acción", "Categoría", "Respuesta", "Palabras clave")
MESSAGE_COLUMNS = ("Clave", "Texto", "Para qué se usa")
REQUEST_COLUMNS = ("Fecha", "Número", "Nombre", "Resumen", "Estado")

MenuAction = Literal["responder", "productos", "persona"]
MENU_ACTIONS: tuple[MenuAction, ...] = ("responder", "productos", "persona")

# Fallback texts for the bot's own messages. They are not business data, so they
# may live in code; the "Mensajes" sheet overrides any of them.
# Keys, default text, and a description shown in the Excel template.
DEFAULT_MESSAGES: dict[str, tuple[str, str]] = {
    "bienvenida": (
        "¡Hola! Gracias por escribir a Bosca.",
        "Saludo al inicio de una conversación (modo menú).",
    ),
    "menu_intro": (
        "Escribe el número de la opción que te interesa:",
        "Texto antes de la lista de opciones del menú.",
    ),
    "no_entiendo": (
        "No tengo información sobre eso. Puedo ayudarte con estas opciones:",
        "Cuando el bot no encuentra una respuesta (modo menú). Va seguido del menú.",
    ),
    "fuera_de_tema": (
        "Solo puedo ayudarte con temas de Bosca. ¿Te interesa alguna caja musical?",
        "Respuesta a temas que no son del negocio (modo IA).",
    ),
    "handoff_cliente": (
        "Gracias por tu mensaje. Pronto te atenderá una persona del equipo de Bosca.",
        "Lo que recibe el cliente cuando se pasa la conversación a una persona.",
    ),
    "sin_precio": (
        "El precio de esta pieza se cotiza según lo que necesites.",
        "Se muestra cuando un producto no tiene precio en la hoja Productos.",
    ),
    "cierre_producto": (
        "Si quieres una cotización o hablar con una persona, escribe {opcion_persona}.",
        "Se agrega después de la información de un producto. {opcion_persona} se reemplaza por el número de la opción de la hoja Menu con acción persona.",
    ),
    "sin_productos": (
        "Por ahora no tengo productos para mostrar en esta categoría.",
        "Cuando una opción del menú no encuentra productos.",
    ),
    "solo_texto": (
        "Por ahora solo puedo leer mensajes de texto. ¿En qué te puedo ayudar?",
        "Respuesta a fotos, audios, stickers y otros mensajes que no son texto.",
    ),
}


class KnowledgeError(ValueError):
    """Raised when the workbook cannot be parsed."""


@dataclass(frozen=True)
class Product:
    name: str
    category: str = ""
    description: str = ""
    customization: str = ""
    price: str = ""
    keywords: tuple[str, ...] = ()


@dataclass(frozen=True)
class FaqEntry:
    question: str
    answer: str
    keywords: tuple[str, ...] = ()


@dataclass(frozen=True)
class MenuOption:
    number: str
    label: str
    action: MenuAction = "responder"
    category: str = ""
    answer: str = ""
    keywords: tuple[str, ...] = ()


@dataclass(frozen=True)
class Knowledge:
    products: tuple[Product, ...] = ()
    faq: tuple[FaqEntry, ...] = ()
    menu: tuple[MenuOption, ...] = ()
    messages: dict[str, str] = field(default_factory=dict)
    loaded: bool = True

    @classmethod
    def empty(cls) -> "Knowledge":
        """Knowledge used before the Excel loads for the first time."""
        return cls(loaded=False)

    def message(self, key: str, **values: str) -> str:
        """Return a bot message, from the Excel if present, else the default."""
        template = self.messages.get(key) or DEFAULT_MESSAGES[key][0]
        for name, value in values.items():
            template = template.replace("{" + name + "}", value)
        return template

    def handoff_option(self) -> MenuOption | None:
        return next((option for option in self.menu if option.action == "persona"), None)

    def products_in(self, category: str) -> tuple[Product, ...]:
        if not category.strip():
            return self.products
        wanted = normalize(category)
        return tuple(product for product in self.products if normalize(product.category) == wanted)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def format_price(value: object) -> str:
    """Numbers become colones (25000 -> ₡25.000). Text is kept as the owner wrote it."""
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        amount = int(value) if float(value).is_integer() else round(float(value), 2)
        return "₡" + f"{amount:,}".replace(",", " ").replace(".", ",").replace(" ", ".")
    return _cell_text(value)


def _split_keywords(value: object) -> tuple[str, ...]:
    parts = (normalize(part) for part in _cell_text(value).replace(";", ",").split(","))
    return tuple(part for part in parts if part)


def _find_sheet(workbook, name: str) -> Worksheet | None:
    wanted = normalize(name)
    for sheet in workbook.worksheets:
        if normalize(sheet.title) == wanted:
            return sheet
    return None


def _rows(sheet: Worksheet | None, columns: tuple[str, ...]) -> list[dict[str, object]]:
    """Read a sheet as dicts keyed by the expected column names.

    The first row holds the headers. Missing optional columns read as empty, and
    fully empty rows are skipped.
    """
    if sheet is None:
        return []
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    positions: dict[str, int] = {}
    normalized_header = [normalize(_cell_text(cell)) for cell in header]
    for column in columns:
        key = normalize(column)
        if key in normalized_header:
            positions[column] = normalized_header.index(key)
    result = []
    for row in rows:
        if all(_cell_text(cell) == "" for cell in row):
            continue
        result.append(
            {column: (row[index] if index < len(row) else None) for column, index in positions.items()}
        )
    return result


def _is_available(value: object) -> bool:
    text = normalize(_cell_text(value))
    return text not in ("no", "n", "0", "false", "falso")


def parse_workbook(data: bytes) -> Knowledge:
    """Parse the workbook bytes into Knowledge. Raises KnowledgeError on a broken file."""
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise KnowledgeError(f"Cannot open the Excel file: {exc}") from exc

    try:
        products = tuple(
            Product(
                name=_cell_text(row.get("Nombre")),
                category=_cell_text(row.get("Categoría")),
                description=_cell_text(row.get("Descripción")),
                customization=_cell_text(row.get("Personalización")),
                price=format_price(row.get("Precio")),
                keywords=_split_keywords(row.get("Palabras clave")),
            )
            for row in _rows(_find_sheet(workbook, SHEET_PRODUCTS), PRODUCT_COLUMNS)
            if _cell_text(row.get("Nombre")) and _is_available(row.get("Disponible"))
        )
        faq = tuple(
            FaqEntry(
                question=_cell_text(row.get("Pregunta")),
                answer=_cell_text(row.get("Respuesta")),
                keywords=_split_keywords(row.get("Palabras clave")),
            )
            for row in _rows(_find_sheet(workbook, SHEET_FAQ), FAQ_COLUMNS)
            if _cell_text(row.get("Respuesta"))
        )
        menu = []
        for row in _rows(_find_sheet(workbook, SHEET_MENU), MENU_COLUMNS):
            number = _cell_text(row.get("Opción"))
            label = _cell_text(row.get("Texto del menú"))
            if not number or not label:
                continue
            action = normalize(_cell_text(row.get("Acción"))) or "responder"
            if action not in MENU_ACTIONS:
                raise KnowledgeError(
                    f"Menu option {number}: unknown action {action!r}, use one of {', '.join(MENU_ACTIONS)}"
                )
            menu.append(
                MenuOption(
                    number=number,
                    label=label,
                    action=cast(MenuAction, action),
                    category=_cell_text(row.get("Categoría")),
                    answer=_cell_text(row.get("Respuesta")),
                    keywords=_split_keywords(row.get("Palabras clave")),
                )
            )
        messages = {}
        for row in _rows(_find_sheet(workbook, SHEET_MESSAGES), MESSAGE_COLUMNS):
            key = normalize(_cell_text(row.get("Clave"))).replace(" ", "_")
            text = _cell_text(row.get("Texto"))
            if key in DEFAULT_MESSAGES and text:
                messages[key] = text
    finally:
        workbook.close()

    return Knowledge(products=products, faq=faq, menu=tuple(menu), messages=messages)


# Formatting shared by both engines, so rule mode and AI mode give the same information.


def describe_product(product: Product, knowledge: Knowledge) -> str:
    lines = [f"*{product.name}*"]
    if product.description:
        lines.append(product.description)
    if product.customization:
        lines.append(f"Personalización: {product.customization}")
    lines.append(f"Precio: {product.price}" if product.price else knowledge.message("sin_precio"))
    return "\n".join(lines)


def list_products(products: tuple[Product, ...]) -> str:
    lines = []
    for product in products:
        line = f"- {product.name}"
        if product.description:
            line += f": {product.description}"
        lines.append(line)
    return "\n".join(lines)


def render_menu(knowledge: Knowledge) -> str:
    return "\n".join(f"{option.number}. {option.label}" for option in knowledge.menu)


def knowledge_as_text(knowledge: Knowledge) -> str:
    """Serialize the knowledge as plain text for the AI prompt."""
    sections = []
    if knowledge.products:
        blocks = []
        for product in knowledge.products:
            block = [f"Producto: {product.name}"]
            if product.category:
                block.append(f"Categoría: {product.category}")
            if product.description:
                block.append(f"Descripción: {product.description}")
            if product.customization:
                block.append(f"Personalización: {product.customization}")
            block.append(f"Precio: {product.price}" if product.price else "Precio: no indicado, se cotiza")
            blocks.append("\n".join(block))
        sections.append("## Productos\n" + "\n\n".join(blocks))
    if knowledge.faq:
        sections.append(
            "## Preguntas frecuentes\n"
            + "\n\n".join(f"Pregunta: {entry.question}\nRespuesta: {entry.answer}" for entry in knowledge.faq)
        )
    answers = [option for option in knowledge.menu if option.action == "responder" and option.answer]
    if answers:
        sections.append(
            "## Información adicional\n"
            + "\n\n".join(f"Tema: {option.label}\n{option.answer}" for option in answers)
        )
    return "\n\n".join(sections) if sections else "(No hay información cargada.)"
