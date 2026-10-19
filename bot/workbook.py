"""Cached knowledge plus serialized writes to the requests sheet."""

import asyncio
import logging
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from .excel_store import ExcelStore
from .knowledge import REQUEST_COLUMNS, SHEET_REQUESTS, Knowledge, parse_workbook
from .text import normalize

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class HandoffRequest:
    created_at: datetime
    number: str
    name: str
    summary: str
    status: str = "Pendiente"


def append_request_row(data: bytes, request: HandoffRequest) -> bytes:
    """Return a copy of the workbook bytes with the request added to the requests sheet.

    Creates the sheet with its headers if it does not exist yet.
    """
    workbook = load_workbook(BytesIO(data))
    sheet = next(
        (ws for ws in workbook.worksheets if normalize(ws.title) == normalize(SHEET_REQUESTS)),
        None,
    )
    if sheet is None:
        sheet = workbook.create_sheet(SHEET_REQUESTS)
        sheet.append(list(REQUEST_COLUMNS))
    sheet.append(
        [
            request.created_at.strftime("%Y-%m-%d %H:%M"),
            # Stored as text so Excel does not turn the number into 5.07E+10.
            request.number,
            request.name,
            request.summary,
            request.status,
        ]
    )
    sheet.cell(row=sheet.max_row, column=2).number_format = "@"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class WorkbookService:
    def __init__(self, store: ExcelStore) -> None:
        self._store = store
        self._knowledge = Knowledge.empty()
        self._write_lock = asyncio.Lock()

    @property
    def knowledge(self) -> Knowledge:
        return self._knowledge

    async def reload(self) -> bool:
        """Download and parse the Excel. Keeps the previous knowledge if it fails."""
        try:
            data = await asyncio.to_thread(self._store.download)
            knowledge = await asyncio.to_thread(parse_workbook, data)
        except Exception:
            logger.exception("Excel reload failed; keeping the previous knowledge")
            return False
        self._knowledge = knowledge
        logger.info(
            "Excel loaded: %d products, %d questions, %d menu options",
            len(knowledge.products),
            len(knowledge.faq),
            len(knowledge.menu),
        )
        return True

    async def add_request(self, request: HandoffRequest) -> bool:
        """Append a row to the requests sheet.

        Writes are serialized, and the file is downloaded right before writing so
        recent edits by the owner are kept.
        """
        async with self._write_lock:
            try:
                data = await asyncio.to_thread(self._store.download)
                updated = await asyncio.to_thread(append_request_row, data, request)
                await asyncio.to_thread(self._store.upload, updated)
            except Exception:
                logger.exception("Could not write the request to the Excel")
                return False
        return True
