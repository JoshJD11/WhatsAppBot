from io import BytesIO

import pytest
from openpyxl import Workbook

from bot.knowledge import (
    DEFAULT_MESSAGES,
    Knowledge,
    KnowledgeError,
    describe_product,
    format_price,
    knowledge_as_text,
    list_products,
    parse_workbook,
    render_menu,
)


def _bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_workbook_reads_available_products_only(knowledge):
    names = [product.name for product in knowledge.products]
    assert names == ["Caja Prueba", "Llavero Prueba"]


def test_parse_workbook_formats_price_and_keywords(knowledge):
    box = knowledge.products[0]
    assert box.price == "₡25.000"
    assert box.keywords == ("caja prueba",)
    assert knowledge.products[1].price == ""


def test_parse_workbook_reads_faq_and_menu(knowledge):
    assert knowledge.faq[0].keywords == ("envio", "envios", "enviar")
    assert [option.action for option in knowledge.menu] == ["productos", "productos", "persona"]


def test_parse_workbook_matches_sheet_and_column_names_loosely():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PRODUCTOS"
    sheet.append(["nombre", "descripcion", "precio"])
    sheet.append(["Pieza", "Texto", "Precio a convenir"])
    knowledge = parse_workbook(_bytes(workbook))
    assert knowledge.products[0].description == "Texto"
    assert knowledge.products[0].price == "Precio a convenir"


def test_parse_workbook_rejects_unknown_menu_action():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Menu"
    sheet.append(["Opción", "Texto del menú", "Acción"])
    sheet.append([1, "Algo", "volar"])
    with pytest.raises(KnowledgeError, match="volar"):
        parse_workbook(_bytes(workbook))


def test_parse_workbook_rejects_non_excel_bytes():
    with pytest.raises(KnowledgeError):
        parse_workbook(b"not an excel file")


def test_parse_workbook_message_overrides_default():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mensajes"
    sheet.append(["Clave", "Texto"])
    sheet.append(["bienvenida", "Hola desde el Excel."])
    sheet.append(["clave_desconocida", "Ignorada."])
    knowledge = parse_workbook(_bytes(workbook))
    assert knowledge.message("bienvenida") == "Hola desde el Excel."
    assert knowledge.message("solo_texto") == DEFAULT_MESSAGES["solo_texto"][0]
    assert "clave_desconocida" not in knowledge.messages


def test_message_replaces_placeholders():
    assert Knowledge().message("cierre_producto", opcion_persona="3").endswith("escribe 3.")


def test_empty_knowledge_is_not_loaded():
    assert Knowledge.empty().loaded is False


@pytest.mark.parametrize(
    "value,expected",
    [(25000, "₡25.000"), (1500000.0, "₡1.500.000"), (99.5, "₡99,5"), ("Desde ₡10.000", "Desde ₡10.000"), (None, "")],
)
def test_format_price(value, expected):
    assert format_price(value) == expected


def test_products_in_filters_by_category_ignoring_accents(knowledge):
    assert [product.name for product in knowledge.products_in("cajas MUSICALES")] == ["Caja Prueba"]
    assert len(knowledge.products_in("")) == 2


def test_handoff_option(knowledge):
    assert knowledge.handoff_option().number == "3"
    assert Knowledge().handoff_option() is None


def test_describe_product_with_and_without_price(knowledge):
    with_price = describe_product(knowledge.products[0], knowledge)
    assert "Precio: ₡25.000" in with_price
    assert "Personalización: Grabado de nombre" in with_price
    without_price = describe_product(knowledge.products[1], knowledge)
    assert DEFAULT_MESSAGES["sin_precio"][0] in without_price


def test_list_products_and_render_menu(knowledge):
    assert list_products(knowledge.products) == "- Caja Prueba: Caja de prueba.\n- Llavero Prueba: Llavero de prueba."
    assert render_menu(knowledge).splitlines()[0] == "1. Cajas musicales"


def test_knowledge_as_text_includes_all_sources(knowledge):
    text = knowledge_as_text(knowledge)
    assert "Producto: Caja Prueba" in text
    assert "Precio: no indicado, se cotiza" in text
    assert "Respuesta de prueba sobre envíos." in text
    assert "Oculto Prueba" not in text


def test_knowledge_as_text_when_empty():
    assert knowledge_as_text(Knowledge()) == "(No hay información cargada.)"
