import asyncio
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from bot.knowledge import REQUEST_COLUMNS
from bot.workbook import HandoffRequest, WorkbookService, append_request_row

from .conftest import MemoryStore

REQUEST = HandoffRequest(
    created_at=datetime(2026, 10, 16, 9, 30),
    number="50688887777",
    name="Cliente",
    summary="Quiere cotizar.",
)


def _rows(data: bytes, sheet: str = "Solicitudes") -> list[tuple]:
    return list(load_workbook(BytesIO(data))[sheet].iter_rows(values_only=True))


def test_append_request_row_adds_row_as_text(workbook_bytes):
    rows = _rows(append_request_row(workbook_bytes, REQUEST))
    assert rows[0] == REQUEST_COLUMNS
    assert rows[1] == ("2026-10-16 09:30", "50688887777", "Cliente", "Quiere cotizar.", "Pendiente")


def test_append_request_row_creates_missing_sheet_and_keeps_others():
    workbook = Workbook()
    workbook.active.title = "Productos"
    workbook.active.append(["Nombre"])
    buffer = BytesIO()
    workbook.save(buffer)
    updated = append_request_row(buffer.getvalue(), REQUEST)
    assert _rows(updated)[0] == REQUEST_COLUMNS
    assert _rows(updated, "Productos") == [("Nombre",)]


def test_reload_loads_knowledge(workbook_bytes):
    service = WorkbookService(MemoryStore(workbook_bytes))
    assert not service.knowledge.loaded
    assert asyncio.run(service.reload()) is True
    assert service.knowledge.loaded
    assert service.knowledge.products[0].name == "Caja Prueba"


def test_reload_failure_keeps_previous_knowledge(workbook_bytes):
    store = MemoryStore(workbook_bytes)
    service = WorkbookService(store)
    asyncio.run(service.reload())
    store.data = b"broken"
    assert asyncio.run(service.reload()) is False
    assert service.knowledge.products[0].name == "Caja Prueba"


def test_add_request_redownloads_and_uploads(workbook_bytes):
    store = MemoryStore(workbook_bytes)
    service = WorkbookService(store)

    async def scenario():
        return await asyncio.gather(service.add_request(REQUEST), service.add_request(REQUEST))

    assert asyncio.run(scenario()) == [True, True]
    assert store.uploads == 2
    assert len(_rows(store.data)) == 3


def test_add_request_returns_false_when_store_fails(workbook_bytes):
    store = MemoryStore(workbook_bytes)
    store.fail = True
    assert asyncio.run(WorkbookService(store).add_request(REQUEST)) is False
