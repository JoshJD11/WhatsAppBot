import asyncio
import itertools
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from bot.bot import owner_notification
from bot.knowledge import DEFAULT_MESSAGES
from bot.main import build_services
from bot.state import BotState, StateFile
from bot.whatsapp import IncomingMessage, OwnerReply
from bot.workbook import HandoffRequest

from .conftest import CUSTOMER, OWNER, PHONE_NUMBER_ID, FakeProvider, FakeSender, MemoryStore, ai_json, make_settings

HANDOFF_TEXT = DEFAULT_MESSAGES["handoff_cliente"][0]


_message_ids = itertools.count(1)


def message(text="hola", sender=CUSTOMER, message_id=None, kind="text", phone_number_id=PHONE_NUMBER_ID):
    return IncomingMessage(
        message_id=message_id or f"wamid.{next(_message_ids)}",
        sender=sender,
        name="Cliente Prueba",
        type=kind,
        text=text,
        phone_number_id=phone_number_id,
    )



class Harness:
    def __init__(self, workbook_bytes, clock, provider=None, state_file=None, **settings_overrides):
        self.sender = FakeSender()
        self.store = MemoryStore(workbook_bytes)
        self.services = build_services(
            make_settings(**settings_overrides), self.sender, self.store, provider, state_file
        )
        # Use the test clock for every time-based rule.
        self.services.state._clock = clock
        asyncio.run(self.services.workbook.reload())

    @property
    def state(self) -> BotState:
        return self.services.state

    def send(self, *messages: IncomingMessage) -> None:
        async def scenario():
            for item in messages:
                await self.services.bot.handle(item)

        asyncio.run(scenario())

    def owner_reply(self, customer=CUSTOMER, message_id=None, phone_number_id=PHONE_NUMBER_ID) -> None:
        reply = OwnerReply(message_id or f"wamid.{next(_message_ids)}", customer, phone_number_id)
        asyncio.run(self.services.bot.handle_owner_reply(reply))

    def to_customer(self) -> list[str]:
        return self.sender.texts_to(CUSTOMER)

    def to_owner(self) -> list[str]:
        return self.sender.texts_to(OWNER)

    def request_rows(self) -> list[tuple]:
        return list(load_workbook(BytesIO(self.store.data))["Solicitudes"].iter_rows(values_only=True))[1:]


@pytest.fixture
def harness(workbook_bytes, clock):
    return Harness(workbook_bytes, clock)


def test_customer_gets_rule_reply(harness):
    harness.send(message("hola"))
    assert harness.to_customer()[0].startswith(DEFAULT_MESSAGES["bienvenida"][0])


def test_duplicate_message_is_answered_once(harness):
    harness.send(message("hola", message_id="same"), message("hola", message_id="same"))
    assert len(harness.to_customer()) == 1


def test_message_for_other_phone_number_id_is_ignored(harness):
    harness.send(message("hola", phone_number_id="other"))
    assert harness.sender.sent == []


def test_non_text_message_gets_text_only_reply(harness):
    harness.send(message("", kind="image"))
    assert harness.to_customer() == [DEFAULT_MESSAGES["solo_texto"][0]]


def test_admin_commands_switch_mode_and_confirm(harness):
    harness.send(message("#bosca ia", sender=OWNER))
    assert harness.state.mode == "ai"
    harness.send(message("#BOSCA MENU", sender=OWNER))
    assert harness.state.mode == "rule"
    assert harness.to_owner() == ["Modo IA activado.", "Modo menú activado."]


def test_other_admin_messages_get_no_reply_but_open_window(harness):
    assert not harness.state.owner_window_open()
    harness.send(message("buenos días", sender=OWNER))
    assert harness.sender.sent == []
    assert harness.state.owner_window_open()


def test_customer_cannot_run_commands(harness):
    harness.send(message("#bosca ia"))
    assert harness.state.mode == "rule"
    assert "Modo IA activado." not in harness.to_customer()


def test_handoff_notifies_owner_logs_request_and_pauses(harness, clock):
    harness.send(message("x", sender=OWNER))  # opens the owner's window
    harness.send(message("3"))

    assert harness.to_customer() == [HANDOFF_TEXT]
    [notification] = harness.to_owner()
    assert "+506 8888 7777" in notification
    assert "Cliente Prueba" in notification
    rows = harness.request_rows()
    assert rows[0][1:3] == ("50688887777", "Cliente Prueba")
    assert "Hablar con una persona" in rows[0][3]
    assert rows[0][4] == "Pendiente"

    clock.advance(minutes=59)
    harness.send(message("¿hola?"))
    assert harness.to_customer() == [HANDOFF_TEXT]

    clock.advance(minutes=1)
    harness.send(message("hola"))
    assert len(harness.to_customer()) == 2


def test_handoff_with_closed_window_still_replies_and_logs(harness):
    harness.send(message("quiero hablar con una persona"))
    assert harness.to_customer() == [HANDOFF_TEXT]
    assert harness.to_owner() == []
    assert len(harness.request_rows()) == 1


def test_handoff_survives_excel_write_failure(harness):
    harness.store.fail = True
    harness.send(message("3"))
    assert harness.to_customer() == [HANDOFF_TEXT]
    assert harness.state.is_paused(CUSTOMER)


def test_rate_limit_silences_spam(workbook_bytes, clock):
    harness = Harness(workbook_bytes, clock, rate_limit_max_messages=2)
    harness.send(message("hola"), message("hola"), message("hola"))
    assert len(harness.to_customer()) == 2


def test_ai_mode_answers_with_ai_and_keeps_history(workbook_bytes, clock):
    provider = FakeProvider(ai_json("answer", "Respuesta IA uno."), ai_json("answer", "Respuesta IA dos."))
    harness = Harness(workbook_bytes, clock, provider=provider, bot_mode="ai")
    harness.send(message("primera"), message("segunda"))
    assert harness.to_customer() == ["Respuesta IA uno.", "Respuesta IA dos."]
    assert [turn.text for turn in provider.calls[1]["history"]] == ["primera", "Respuesta IA uno."]


def test_ai_failure_falls_back_to_rule_mode_for_that_message_only(workbook_bytes, clock):
    provider = FakeProvider(RuntimeError("429 quota"), ai_json("answer", "Vuelve la IA."))
    harness = Harness(workbook_bytes, clock, provider=provider, bot_mode="ai")
    harness.send(message("¿hacen envíos?"))
    assert harness.to_customer() == ["Respuesta de prueba sobre envíos."]
    assert harness.state.mode == "ai"
    harness.send(message("otra"))
    assert harness.to_customer()[-1] == "Vuelve la IA."


def test_ai_mode_without_key_always_uses_rules(workbook_bytes, clock):
    harness = Harness(workbook_bytes, clock, provider=None, bot_mode="ai")
    harness.send(message("1"))
    assert "Caja Prueba" in harness.to_customer()[0]


def test_ai_handoff_uses_model_summary(workbook_bytes, clock):
    provider = FakeProvider(ai_json("handoff", "", "Quiere una caja con grabado."))
    harness = Harness(workbook_bytes, clock, provider=provider, bot_mode="ai")
    harness.send(message("x", sender=OWNER), message("quiero encargar una caja"))
    assert harness.to_customer() == [HANDOFF_TEXT]
    assert "Quiere una caja con grabado." in harness.to_owner()[0]
    assert harness.request_rows()[0][3] == "Quiere una caja con grabado."


def test_outgoing_ai_text_is_sanitized(workbook_bytes, clock):
    provider = FakeProvider(ai_json("answer", "Claro 😊 -- con gusto"))
    harness = Harness(workbook_bytes, clock, provider=provider, bot_mode="ai")
    harness.send(message("hola"))
    assert harness.to_customer() == ["Claro - con gusto"]


def test_failed_send_is_not_remembered(workbook_bytes, clock):
    provider = FakeProvider(ai_json("answer", "uno"), ai_json("answer", "dos"))
    harness = Harness(workbook_bytes, clock, provider=provider, bot_mode="ai")
    harness.sender.succeed = False
    harness.send(message("primera"), message("segunda"))
    assert [turn.role for turn in provider.calls[1]["history"]] == ["customer"]


def test_missing_owner_number_skips_notification(workbook_bytes, clock):
    harness = Harness(workbook_bytes, clock, owner_number="", admin_numbers=frozenset())
    harness.send(message("3"))
    assert harness.to_customer() == [HANDOFF_TEXT]
    assert len(harness.sender.sent) == 1


def test_owner_notification_text():
    request = HandoffRequest(datetime(2026, 1, 1), "50688887777", "", "Resumen.")
    text = owner_notification(request, 60)
    assert "Nombre" not in text
    assert "Resumen: Resumen." in text
    assert "60 minutos" in text
    assert "--" not in text


def test_owner_reply_pauses_the_bot_for_that_customer(harness, clock):
    harness.owner_reply()
    harness.send(message("hola"))
    assert harness.to_customer() == []
    clock.advance(minutes=60)
    harness.send(message("hola"))
    assert len(harness.to_customer()) == 1


def test_each_owner_reply_extends_the_pause(harness, clock):
    harness.owner_reply()
    clock.advance(minutes=50)
    harness.owner_reply()
    clock.advance(minutes=50)
    harness.send(message("hola"))
    assert harness.to_customer() == []


def test_owner_reply_clears_ai_history(workbook_bytes, clock):
    provider = FakeProvider(ai_json("answer", "uno"), ai_json("answer", "dos"))
    harness = Harness(workbook_bytes, clock, provider=provider, bot_mode="ai")
    harness.send(message("primera"))
    harness.owner_reply()
    clock.advance(minutes=60)
    harness.send(message("segunda"))
    assert provider.calls[1]["history"] == []


def test_owner_reply_for_other_phone_number_id_is_ignored(harness):
    harness.owner_reply(phone_number_id="other")
    assert not harness.state.is_paused(CUSTOMER)


def test_duplicate_owner_reply_is_ignored(harness, clock):
    harness.owner_reply(message_id="echo-1")
    clock.advance(minutes=50)
    harness.owner_reply(message_id="echo-1")
    clock.advance(minutes=10)
    assert not harness.state.is_paused(CUSTOMER)


def test_state_is_saved_on_mode_switch_owner_message_handoff_and_owner_reply(workbook_bytes, clock, tmp_path):
    state_file = StateFile(str(tmp_path / "state.json"))
    harness = Harness(workbook_bytes, clock, state_file=state_file)

    harness.send(message("#bosca ia", sender=OWNER))
    saved = state_file.load()
    assert saved["mode"] == "ai"
    assert saved["owner_last_message"] is not None

    harness.send(message("#bosca menu", sender=OWNER), message("3"))
    assert CUSTOMER in state_file.load()["paused_until"]

    harness.owner_reply(customer="50611112222")
    assert "50611112222" in state_file.load()["paused_until"]


def test_state_save_failure_does_not_stop_the_bot(workbook_bytes, clock, tmp_path):
    class BrokenStateFile(StateFile):
        def save(self, snapshot):
            raise OSError("disk full")

    harness = Harness(workbook_bytes, clock, state_file=BrokenStateFile(str(tmp_path / "s.json")))
    harness.send(message("#bosca ia", sender=OWNER))
    assert harness.state.mode == "ai"
    assert harness.to_owner() == ["Modo IA activado."]

